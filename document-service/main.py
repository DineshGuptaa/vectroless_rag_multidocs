from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import Optional
import pypdf
import httpx
import os
import uuid
import docx

app = FastAPI(title="Document Service", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

class ProcessedDocument(BaseModel):
    filename: str
    page_count: int
    total_chars: int
    estimated_tokens: int
    pages: list[dict]

@app.get("/")
async def root():
    return {"service": "Document Service", "status": "running"}

@app.get("/health")
async def health():
    return {
        "status": "healthy",
        "service": "document-service",
        "port": 8001,
        "version": "1.0.0"
    }

SUPPORTED_FORMATS = {'.pdf', '.docx', '.txt', '.md'}


def extract_text_from_file(file_path: str) -> tuple[list[dict], int]:
    """Extract text from a file based on its extension.

    Returns:
        tuple: (pages list, total_chars)
    """
    ext = os.path.splitext(file_path)[1].lower()

    if ext == '.pdf':
        reader = pypdf.PdfReader(file_path)
        pages = []
        total_chars = 0
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            char_count = len(text)
            total_chars += char_count
            pages.append({
                "page_number": i + 1,
                "text": text,
                "char_count": char_count,
                "estimated_tokens": int(char_count / 4)
            })
        return pages, total_chars

    elif ext == '.docx':
        doc = docx.Document(file_path)
        full_text = '\n'.join(para.text for para in doc.paragraphs)
        char_count = len(full_text)
        total_chars = char_count
        pages = [{
            "page_number": 1,
            "text": full_text,
            "char_count": char_count,
            "estimated_tokens": int(char_count / 4)
        }]
        return pages, total_chars

    elif ext in ('.txt', '.md'):
        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            full_text = f.read()
        char_count = len(full_text)
        total_chars = char_count
        pages = [{
            "page_number": 1,
            "text": full_text,
            "char_count": char_count,
            "estimated_tokens": int(char_count / 4)
        }]
        return pages, total_chars

    else:
        raise ValueError(f"Unsupported file format: {ext}")


@app.post("/upload")
async def upload_document(file: UploadFile = File(...)):
    """Upload and process document (PDF, DOCX, TXT, MD)"""
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in SUPPORTED_FORMATS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file format: {ext}. Supported: {', '.join(sorted(SUPPORTED_FORMATS))}"
        )

    # Save file
    file_id = str(uuid.uuid4())
    file_path = os.path.join(UPLOAD_DIR, f"{file_id}{ext}")

    with open(file_path, "wb") as f:
        content = await file.read()
        f.write(content)

    try:
        # Extract text based on format
        pages, total_chars = extract_text_from_file(file_path)
        page_count = len(pages)
        estimated_tokens = int(total_chars / 4)

        # Store document metadata in storage service
        async with httpx.AsyncClient() as client:
            storage_response = await client.post(
                "http://storage-service:8005/documents",
                json={
                    "filename": file.filename,
                    "file_path": file_path,
                    "size": len(content),
                    "page_count": page_count,
                    "status": "uploaded"
                }
            )

            if storage_response.status_code != 200:
                raise HTTPException(status_code=500, detail="Failed to store document metadata")

            doc_data = storage_response.json()
            doc_id = doc_data["id"]

        return {
            "doc_id": doc_id,
            "filename": file.filename,
            "page_count": page_count,
            "total_chars": total_chars,
            "estimated_tokens": estimated_tokens,
            "status": "processed",
            "message": "Document uploaded and processed successfully"
        }

    except Exception as e:
        # Clean up file on error
        if os.path.exists(file_path):
            os.remove(file_path)
        raise HTTPException(status_code=500, detail=f"Error processing document: {str(e)}")

@app.get("/documents/{doc_id}/extract")
async def extract_document_text(doc_id: int):
    """Extract full text from document"""
    async with httpx.AsyncClient() as client:
        # Get document from storage
        response = await client.get(f"http://storage-service:8005/documents/{doc_id}")

        if response.status_code != 200:
            raise HTTPException(status_code=404, detail="Document not found")

        doc_data = response.json()
        file_path = doc_data["file_path"]

        if not os.path.exists(file_path):
            raise HTTPException(status_code=404, detail="Document file not found")

        # Extract text
        pages, _ = extract_text_from_file(file_path)

        return {
            "doc_id": doc_id,
            "filename": doc_data["filename"],
            "page_count": len(pages),
            "pages": pages
        }

@app.get("/documents/{doc_id}/pages/{page_num}")
async def get_page_text(doc_id: int, page_num: int):
    """Get text from specific page"""
    async with httpx.AsyncClient() as client:
        response = await client.get(f"http://storage-service:8005/documents/{doc_id}")

        if response.status_code != 200:
            raise HTTPException(status_code=404, detail="Document not found")

        doc_data = response.json()
        file_path = doc_data["file_path"]

        if not os.path.exists(file_path):
            raise HTTPException(status_code=404, detail="Document file not found")

        pages, _ = extract_text_from_file(file_path)

        if page_num < 1 or page_num > len(pages):
            raise HTTPException(status_code=400, detail="Invalid page number")

        page = pages[page_num - 1]

        return {
            "doc_id": doc_id,
            "page_number": page_num,
            "text": page["text"],
            "char_count": page["char_count"]
        }
